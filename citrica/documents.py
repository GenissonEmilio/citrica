"""Bundled document conversions. No Office installation or network is required."""
import html
import threading
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from citrica.core import ConversionError, check_cancel


def pdf_to_docx(source: Path, target: Path):
    """Called in a cancellable child process by the shared conversion engine."""
    import pymupdf
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ConversionError('PDF protegido por senha. Desbloqueie uma cópia antes de converter.')
        if not len(document):
            raise ConversionError('O PDF não contém páginas.')
    from pdf2docx import Converter
    converter = Converter(str(source))
    try:
        converter.convert(str(target), multi_processing=False, ignore_page_error=False,
                          raw_exceptions=True)
    finally:
        converter.close()


def basic_document(source: Path, target: Path, cancel: threading.Event) -> bool:
    """Fallback for DOCX PDF/TXT and TXT DOCX when LibreOffice is absent.

    This preserves editable text, tables and embedded images; advanced Word
    layout/features are rendered more accurately by the optional LibreOffice.
    """
    from docx import Document
    if source.suffix.lower() == '.txt' and target.suffix == '.docx':
        result = Document()
        result.core_properties.author = ''
        result.core_properties.last_modified_by = ''
        for line in source.read_text(encoding='utf-8-sig').splitlines():
            check_cancel(cancel)
            result.add_paragraph(line)
        result.save(target)
        return True
    if source.suffix.lower() != '.docx' or target.suffix not in {'.pdf', '.txt'}:
        return False
    with zipfile.ZipFile(source) as archive:
        if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 256*1024**2:
            raise ConversionError('O DOCX excede o limite de 256 MiB descompactados ou 10.000 entradas.')
    document = Document(source)
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    if target.suffix == '.txt':
        def lines(container):
            for block in container.iter_inner_content():
                check_cancel(cancel)
                if isinstance(block, Paragraph):
                    yield block.text
                else:
                    for row in block.rows:
                        yield '\t'.join('\n'.join(lines(cell)) for cell in row.cells)
        target.write_text('\n'.join(lines(document)) + '\n', encoding='utf-8')
        return True
    import pymupdf
    from docx.oxml.ns import qn
    from docx.text.hyperlink import Hyperlink
    from PIL import Image
    with TemporaryDirectory(prefix='citrica-docx-') as folder:
        images = {}
        def paragraph(block):
            chunks = []
            runs = (run for part in block.iter_inner_content()
                    for run in (part.runs if isinstance(part, Hyperlink) else [part]))
            for run in runs:
                check_cancel(cancel)
                text = html.escape(run.text).replace('\n', '<br/>').replace('\t', '    ')
                if run.bold:
                    text = '<b>' + text + '</b>'
                if run.italic:
                    text = '<i>' + text + '</i>'
                if run.font.size:
                    text = f'<span style="font-size:{run.font.size.pt}pt">{text}</span>'
                chunks.append(text)
                for blip in run._r.iter(qn('a:blip')):
                    relationship = blip.get(qn('r:embed'))
                    if not relationship:
                        raise ConversionError('O DOCX contém imagem externa; incorpore-a antes de converter.')
                    if relationship not in images:
                        blob = document.part.related_parts[relationship].blob
                        import io
                        with Image.open(io.BytesIO(blob)) as image:
                            image.load()
                            name = f'image-{len(images)}.png'
                            image.convert('RGBA').save(Path(folder)/name)
                            images[relationship] = (name, image.width, image.height)
                    name, width, height = images[relationship]
                    scale = min(1., 450/max(1, width), 600/max(1, height))
                    chunks.append(f'<img src="{name}" width="{width*scale}" height="{height*scale}"/>')
            tag = 'p'
            if block.style and block.style.name.startswith('Heading '):
                level = block.style.name.rsplit(' ', 1)[-1]
                if level.isdigit():
                    tag = 'h' + str(min(6, max(1, int(level))))
            alignment = {1: 'center', 2: 'right', 3: 'justify'}.get(block.alignment, 'left')
            return f'<{tag} style="text-align:{alignment}">' + (''.join(chunks) or '<br/>') + f'</{tag}>'
        def content(container, depth=0):
            if depth > 16:
                raise ConversionError('O DOCX contém tabelas aninhadas em excesso.')
            chunks = []
            for block in container.iter_inner_content():
                check_cancel(cancel)
                if isinstance(block, Paragraph):
                    chunks.append(paragraph(block))
                elif isinstance(block, Table):
                    chunks.append('<table>')
                    for row in block.rows:
                        chunks.append('<tr>')
                        for cell in row.cells:
                            chunks.append('<td>' + content(cell, depth+1) + '</td>')
                        chunks.append('</tr>')
                    chunks.append('</table>')
            return ''.join(chunks)
        story = pymupdf.Story(html=content(document), archive=pymupdf.Archive(folder),
                              user_css='body{font-family:sans-serif;font-size:11pt;} '
                              'table{border-collapse:collapse;width:100%;}td{border:1px solid #777;padding:4pt;}')
        section = document.sections[0]
        def pt(value, default):
            return value.pt if value is not None else default
        width, height = pt(section.page_width, 612), pt(section.page_height, 792)
        area = pymupdf.Rect(pt(section.left_margin, 72), pt(section.top_margin, 72),
                            width-pt(section.right_margin, 72), height-pt(section.bottom_margin, 72))
        if area.width < 20 or area.height < 20:
            raise ConversionError('Margens ou tamanho de página inválidos no DOCX.')
        writer = pymupdf.DocumentWriter(str(target))
        try:
            more = True
            while more:
                check_cancel(cancel)
                device = writer.begin_page(pymupdf.Rect(0, 0, width, height))
                more, _ = story.place(area)
                story.draw(device)
                writer.end_page()
        finally:
            writer.close()
    return True
