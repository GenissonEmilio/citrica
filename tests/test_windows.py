"""Shared-engine and Qt behavior tests for the Windows wheel; Win32 tests run on Windows CI."""
import json
import math
import os
import time
import unicodedata

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest
from PIL import Image
from PySide6.QtCore import QMimeData, QObject, QPoint, QPointF, Qt, QTimer, QUrl, Signal
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication

from citrica.gui import STYLE
from citrica.windows import config
from citrica.windows.app import Desktop
from citrica.windows.wheel import CENTER, LABEL_RADIUS, SIZE, hit_test, offered_paths


class FakeNative(QObject):
    toggle = Signal()
    mode = Signal(str)
    released = Signal()
    key = Signal(str)

    def __init__(self, app, settings):
        super().__init__()
        self.settings, self.visible, self.registered = settings, False, None

    def register(self, text):
        config.parse_shortcut(text)
        if text == 'Ctrl+Alt+X':
            raise ValueError('Atalho em uso')
        self.registered = text

    def reset_drag(self):
        pass

    def left_down(self):
        return False

    def close(self):
        pass


def spin_until(predicate, timeout=5000):
    deadline = time.monotonic() + timeout / 1000
    while time.monotonic() < deadline:
        QApplication.processEvents()
        if predicate():
            return
        # Yield Python's GIL explicitly as well as pumping Qt. This matters for
        # Python QThread workers on the Windows Qt test backend.
        time.sleep(.01)
    controller = getattr(QApplication.instance(), '_windows_test_controller', None)
    raise AssertionError({'state': controller.status() if controller else None,
                          'jobs': [(j.isRunning(), j.isFinished()) for j in controller.jobs] if controller else []})


@pytest.fixture
def desktop(tmp_path, monkeypatch):
    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(STYLE)
    monkeypatch.setenv('APPDATA', str(tmp_path / 'settings'))
    monkeypatch.setenv('XDG_CACHE_HOME', str(tmp_path / 'cache'))
    controller = Desktop(app, FakeNative, listen=False)
    app._windows_test_controller = controller
    # Unit tests must not let an unrelated real cursor override synthetic drops.
    controller.pointer_timer.stop()
    yield controller
    controller.hide_wheel()
    for job in list(controller.jobs):
        job.cancel.set()
        job.wait(5000)
    QApplication.processEvents()
    controller.close()
    app._windows_test_controller = None


def image_file(tmp_path):
    path = tmp_path / 'fóto com espaços.png'
    Image.new('RGB', (320, 200), 'orange').save(path)
    return path


def point_for(desktop, name):
    wheel = desktop.wheel
    index = next(i for i, item in enumerate(wheel.items) if item['id'] == name)
    angle = -math.pi / 2 + index * math.tau / len(wheel.items)
    return QPoint(int(CENTER + LABEL_RADIUS * math.cos(angle)),
                  int(CENTER + LABEL_RADIUS * math.sin(angle)))


def assert_empty(desktop):
    state = desktop.status()
    assert state['paths'] == [] and state['items'] == []
    assert state['selection'] is None and state['selected'] == -1 and not state['cachedSelection']
    assert desktop.wheel._preview.isNull()


def test_windows_hotkey_validation_and_atomic_preferences(tmp_path, monkeypatch):
    monkeypatch.setenv('APPDATA', str(tmp_path))
    assert config.parse_shortcut('Ctrl+Alt+C') == (3, ord('C'))
    assert config.parse_shortcut('Win+Shift+F8') == (12, 0x77)
    for text in ['Shift+C', 'Ctrl+Ctrl+A', 'Win+Super+C', 'C', 'Ctrl+F25', 'Alt+Unknown', '', None]:
        with pytest.raises(ValueError):
            config.parse_shortcut(text)
    data = {'shortcut': 'Win+Shift+F8', 'animations': False, 'drag_key': 'Shift', 'wheel_theme': 'classic'}
    config.save(data)
    assert config.load() == data
    file = config.config_dir() / 'settings.json'
    file.write_text('[invalid', encoding='utf-8')
    assert config.load() == config.DEFAULTS
    file.write_text('[]', encoding='utf-8')
    assert config.load() == config.DEFAULTS


def test_windows_ring_geometry_and_transparent_center(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.show_wheel([source])
    spin_until(lambda: desktop.wheel.items and desktop.wheel.progress == 1)
    for i in range(len(desktop.wheel.items)):
        angle = -math.pi / 2 + i * math.tau / len(desktop.wheel.items)
        x, y = CENTER + LABEL_RADIUS * math.cos(angle), CENTER + LABEL_RADIUS * math.sin(angle)
        assert hit_test(x, y, len(desktop.wheel.items)) == i
        assert desktop.wheel._geometry[i].contains(QPointF(x, y))
    assert hit_test(CENTER, CENTER, 8) == -1
    assert hit_test(0, 0, 8) == -1
    snapshot = desktop.wheel.grab().toImage()
    assert snapshot.pixelColor(1, 1).alpha() <= 2, 'Overlay must not have an opaque rectangular background'
    assert desktop.wheel.width() == SIZE


def test_windows_drop_forces_copy_and_reopen_is_empty(desktop, tmp_path):
    source = image_file(tmp_path)
    before = source.read_bytes()
    desktop.show_wheel([])
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(source))])
    enter = QDragEnterEvent(QPoint(CENTER, CENTER), Qt.DropAction.CopyAction | Qt.DropAction.MoveAction,
                            mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, enter)
    assert enter.isAccepted() and enter.dropAction() == Qt.DropAction.CopyAction
    spin_until(lambda: any(i['id'] == 'jpg' for i in desktop.wheel.items))
    point = point_for(desktop, 'jpg')
    desktop.wheel.select(hit_test(point.x(), point.y(), len(desktop.wheel.items)))
    drop = QDropEvent(QPointF(point), Qt.DropAction.CopyAction | Qt.DropAction.MoveAction, mime,
                     Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, drop)
    assert drop.isAccepted() and drop.dropAction() == Qt.DropAction.CopyAction
    assert_empty(desktop)  # Before the conversion worker or close animation ends.
    spin_until(lambda: not desktop.busy)
    assert source.read_bytes() == before
    assert Image.open(next(tmp_path.glob('*.jpg'))).size == (320, 200)
    desktop.toggle()
    assert desktop.active
    assert_empty(desktop)


def test_windows_error_escape_and_late_reply_do_not_restore_file(desktop, tmp_path):
    source = tmp_path / 'bad.png'
    source.write_bytes(b'not an image')
    desktop.show_wheel([source])
    spin_until(lambda: any(i['id'] == 'jpg' for i in desktop.wheel.items))
    desktop.apply(next(i for i in desktop.wheel.items if i['id'] == 'jpg'))
    spin_until(lambda: not desktop.busy)
    assert desktop.counts['errors'] == 1
    desktop.show_wheel([])
    assert_empty(desktop)
    desktop.hide_wheel()
    source = image_file(tmp_path)
    desktop.show_wheel([source])
    desktop.hide_wheel()  # The capability worker is deliberately still pending.
    desktop.show_wheel([])
    spin_until(lambda: not desktop.jobs)
    assert_empty(desktop)
    desktop.wheel.keyboard('Escape', QPoint())
    assert not desktop.active
    assert_empty(desktop)


def test_windows_tools_actions_pages_keyboard_and_conflict(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.show_wheel([source], 'tools')
    spin_until(lambda: desktop.wheel.items)
    assert any(i['id'] == 'crop' for i in desktop.wheel.items)
    more = next(i for i in desktop.wheel.items if i['kind'] == 'page')
    desktop.apply(more)
    desktop.release()
    QTest.qWait(300)
    assert desktop.active and desktop.paths == [source]
    assert any(i['id'] == 'redact_faces' for i in desktop.wheel.items)
    desktop.drag_mode('actions')
    copy = next(i for i in desktop.wheel.items if i['id'] == 'copy_path')
    desktop.apply(copy)
    spin_until(lambda: desktop.app.clipboard().text() == str(source))
    assert_empty(desktop)
    desktop.show_wheel([source])
    spin_until(lambda: desktop.wheel.items)
    desktop.wheel.keyboard('Right', QPoint(10, 10))
    desktop.wheel.pointer(QPoint(10, 10))
    assert desktop.wheel.selected == 0
    with pytest.raises(ValueError, match='em uso'):
        desktop.apply_settings({'shortcut': 'Ctrl+Alt+X', 'animations': False, 'drag_key': 'Win'})
    assert desktop.settings == config.DEFAULTS


def test_windows_mime_rejects_remote_and_missing_files(tmp_path):
    mime = QMimeData()
    source = image_file(tmp_path)
    mime.setUrls([QUrl.fromLocalFile(str(source)), QUrl('https://example.com/file.png')])
    assert offered_paths(mime) == []
    mime.setUrls([QUrl.fromLocalFile(str(tmp_path / 'missing.png'))])
    assert offered_paths(mime) == []
    mime.setUrls([QUrl.fromLocalFile(str(source))] * 2)
    assert offered_paths(mime) == [source]


def test_windows_ipc_rejects_invalid_request(desktop):
    with pytest.raises(ValueError):
        desktop.dispatch({'op': 'show', 'paths': 'not-a-list'})
    with pytest.raises(ValueError):
        desktop.dispatch({'op': 'execute-arbitrary-command'})
    assert json.loads(json.dumps(desktop.dispatch({'op': 'status'})))['version']


def test_windows_drag_probe_reveals_only_for_an_actual_file(desktop, tmp_path):
    desktop.drag_mode('convert')
    assert desktop.armed and not desktop.active and desktop.wheel.probing
    assert desktop.wheel.grab().toImage().pixelColor(CENTER, 30).alpha() <= 2
    text = QMimeData()
    text.setText('This is a text selection, not a file')
    enter = QDragEnterEvent(QPoint(CENTER, CENTER), Qt.DropAction.CopyAction,
                            text, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, enter)
    assert not enter.isAccepted() and not desktop.active and not desktop.armed
    desktop.drag_mode('convert')
    source = image_file(tmp_path)
    files = QMimeData()
    files.setUrls([QUrl.fromLocalFile(str(source))])
    enter = QDragEnterEvent(QPoint(CENTER, CENTER), Qt.DropAction.CopyAction,
                            files, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, enter)
    spin_until(lambda: desktop.wheel.items)
    assert desktop.active and not desktop.armed and not desktop.wheel.probing


def test_windows_copy_keeps_resident_alive_after_clipboard_changes(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.show_wheel([source], 'actions')
    spin_until(lambda: desktop.wheel.items)
    copy = next(i for i in desktop.wheel.items if i['id'] == 'copy')
    desktop.apply(copy)
    spin_until(lambda: desktop.app.clipboard().mimeData().urls() == [QUrl.fromLocalFile(str(source))])
    assert desktop.app.clipboard().mimeData().urls() == [QUrl.fromLocalFile(str(source))]
    desktop.app.clipboard().setText('The user copied something else')
    previous = desktop.app.quitOnLastWindowClosed()
    desktop.app.setQuitOnLastWindowClosed(False)
    started = time.monotonic()
    QTimer.singleShot(1300, desktop.app.quit)
    desktop.app.exec()
    desktop.app.setQuitOnLastWindowClosed(previous)
    assert time.monotonic() - started >= 1.2, 'Clipboard changes must not terminate the resident app'


def test_windows_explorer_filenamew_offer(tmp_path):
    source = image_file(tmp_path)
    mime = QMimeData()
    name = 'application/x-qt-windows-mime;value="FileNameW"'
    mime.setData(name, (str(source)+'\0\0').encode('utf-16-le'))
    assert not mime.hasUrls()
    assert offered_paths(mime) == [source]
    mime.setData(name, b'odd')
    assert offered_paths(mime) == []
    import struct
    drop = QMimeData()
    data = struct.pack('<IiiII', 20, 0, 0, 0, 1) + (str(source)+'\0\0').encode('utf-16-le')
    drop.setData('application/x-qt-windows-mime;value="CF_HDROP"', data)
    assert offered_paths(drop) == [source]


def test_windows_repeated_probe_tracks_the_new_pointer(desktop, monkeypatch):
    point = QPoint(200, 200)
    monkeypatch.setattr('citrica.windows.app.QCursor.pos', lambda: QPoint(point))
    desktop.drag_mode('actions')
    first = desktop.wheel.pos()
    point.setX(500)
    desktop.drag_mode('actions')
    assert desktop.armed and desktop.wheel.probing
    assert desktop.wheel.pos() != first


def test_windows_explorer_preview_populates_before_ole_but_requires_real_offer(desktop, tmp_path, monkeypatch):
    source = image_file(tmp_path)
    monkeypatch.setattr('citrica.windows.explorer.preview_files', lambda hwnd: [source])
    monkeypatch.setattr(desktop.native, 'left_down', lambda: True)
    desktop.prepare_drag(123)
    spin_until(lambda: desktop.drag_candidates)
    desktop.drag_mode('actions')
    spin_until(lambda: desktop.wheel.items)
    assert desktop.active and not desktop.drag_verified
    before = desktop.app.clipboard().text()
    copy = next(i for i in desktop.wheel.items if i['id'] == 'copy_path')
    desktop.apply(copy)
    assert desktop.active and desktop.app.clipboard().text() == before
    desktop.offer([source])
    assert desktop.drag_verified
    desktop.apply(copy)
    spin_until(lambda: desktop.app.clipboard().text() == str(source))
    assert_empty(desktop)


def test_windows_fast_drop_waits_for_slow_palette(desktop, tmp_path, monkeypatch):
    from citrica.desktop import request
    def delayed(payload):
        time.sleep(.7)
        return request(payload)
    monkeypatch.setattr('citrica.windows.app.request', delayed)
    source = image_file(tmp_path)
    desktop.show_wheel([])
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(source))])
    enter = QDragEnterEvent(QPoint(CENTER, CENTER), Qt.DropAction.CopyAction,
                            mime, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, enter)
    angle = -math.pi/2 + math.tau/8  # JPG is sector 1 on image page 1.
    point = QPointF(CENTER+LABEL_RADIUS*math.cos(angle), CENTER+LABEL_RADIUS*math.sin(angle))
    drop = QDropEvent(point, Qt.DropAction.CopyAction, mime,
                     Qt.MouseButton.LeftButton, Qt.KeyboardModifier.ShiftModifier)
    QApplication.sendEvent(desktop.wheel, drop)
    assert drop.isAccepted() and desktop.pending_drop and not desktop.dragging
    desktop.release()
    spin_until(lambda: desktop.counts['converted'] == 1 and not desktop.busy)
    assert Image.open(next(tmp_path.glob('*.jpg'))).size == (320, 200)
    assert_empty(desktop)


def test_windows_docx_and_pdf_options_without_libreoffice(desktop, tmp_path, monkeypatch):
    from docx import Document
    import pymupdf
    monkeypatch.setenv('CITRICA_OFFICE_ENGINE', 'builtin')
    source = tmp_path/'document.docx'
    doc = Document()
    doc.add_paragraph('Citrica without LibreOffice')
    doc.save(source)
    desktop.show_wheel([source])
    spin_until(lambda: desktop.wheel.items)
    assert {i['id'] for i in desktop.wheel.items} == {'pdf', 'txt'}
    desktop.apply(next(i for i in desktop.wheel.items if i['id'] == 'pdf'))
    spin_until(lambda: desktop.counts['converted'] == 1 and not desktop.busy)
    pdf = next(tmp_path.glob('*.pdf'))
    with pymupdf.open(pdf) as result:
        assert 'Citrica without LibreOffice' in unicodedata.normalize('NFKC', result[0].get_text())
    desktop.show_wheel([pdf])
    spin_until(lambda: desktop.wheel.items)
    assert any(i['id'] == 'docx' for i in desktop.wheel.items)
    desktop.apply(next(i for i in desktop.wheel.items if i['id'] == 'docx'))
    spin_until(lambda: desktop.counts['converted'] == 2 and not desktop.busy)
    assert any('Citrica without LibreOffice' in unicodedata.normalize('NFKC', p.text)
               for p in Document(tmp_path/'document (1).docx').paragraphs)


def test_drop_in_empty_center_dismisses_and_reopening_is_empty(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.show_wheel([source], dragging=True)
    spin_until(lambda: desktop.wheel.items)
    desktop.drop([source], QPoint(CENTER, CENTER))
    spin_until(lambda: not desktop.active)
    assert_empty(desktop)
    assert desktop.counts == {'converted': 0, 'errors': 0}
    desktop.show_wheel([])
    QTest.qWait(300)
    assert not desktop.paths and not desktop.wheel.info


def test_release_outside_survives_late_offer_generation_change(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.drag_mode('actions')
    desktop.release()
    # Explorer can deliver the OLE offer after the native button-up. This
    # changes the palette generation, but still belongs to the same drag.
    desktop.offer([source])
    spin_until(lambda: not desktop.active and not desktop.armed)
    assert_empty(desktop)


def test_button_state_fallback_dismisses_when_native_release_is_missed(desktop, tmp_path):
    source = image_file(tmp_path)
    desktop.show_wheel([source], dragging=True)
    desktop.pointer_timer.start(20)
    spin_until(lambda: not desktop.active)
    assert_empty(desktop)


def test_stale_release_does_not_dismiss_a_new_drag(desktop, tmp_path, monkeypatch):
    source = image_file(tmp_path)
    desktop.show_wheel([source], dragging=True)
    desktop.release()
    desktop.hide_wheel()
    monkeypatch.setattr(desktop.native, 'left_down', lambda: True)
    desktop.show_wheel([source], dragging=True)
    QTest.qWait(350)
    assert desktop.active and desktop.dragging
