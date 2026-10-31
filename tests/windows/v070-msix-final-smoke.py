"""Final installed-QA MSIX smoke; actual PFN and Popen ownership guards."""

import argparse
import ctypes
from ctypes import wintypes
import json
import os
from pathlib import Path
import shutil
import subprocess
from docx import Document
from v070_editor_cases import Driver, digest, engine, duration, play_and_pause, playback_seconds

import pymupdf

FAMILY = "Emilians.Ctrica.UnsignedQA_p9dw1q8pw0bz6"
p = argparse.ArgumentParser()
p.add_argument("--gui", type=Path, required=True)
p.add_argument("--state", type=Path, required=True)
p.add_argument("--oracle-ffprobe", type=Path, required=True)
a = p.parse_args()
assert not a.state.exists()
a.state.mkdir()
assert digest(a.gui) == "5cf40dd9550b41becfe6f902b4259a64e777d2a61370d5eb775e9942c8d89688"
assert "WindowsApps" in str(a.gui) and a.gui.parts[-3].startswith("Emilians.Ctrica.UnsignedQA_")
assert digest(a.oracle_ffprobe) == digest(engine(a.gui, "ffprobe"))
driver = Driver(a.gui, Path(__file__).with_name("editor_uia.ps1"), a.state)
report = {
    "result": "running",
    "scope": "installed unsigned QA MSIX only, final frozen GUI",
    "gui_sha256": digest(a.gui),
    "package_family": FAMILY,
    "duration_oracle": "SHA-identical ffprobe from the validated portable; not a direct WindowsApps helper launch",
    "oracle_ffprobe_sha256": digest(a.oracle_ffprobe),
    "cases": {},
}


def medium_integrity(pid):
    kernel = driver.kernel
    adv = ctypes.WinDLL("advapi32", use_last_error=True)
    adv.OpenProcessToken.argtypes = [wintypes.HANDLE, wintypes.DWORD, ctypes.POINTER(wintypes.HANDLE)]
    adv.GetTokenInformation.argtypes = [
        wintypes.HANDLE,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        ctypes.POINTER(wintypes.DWORD),
    ]
    adv.GetSidSubAuthorityCount.argtypes = [wintypes.LPVOID]
    adv.GetSidSubAuthorityCount.restype = ctypes.POINTER(ctypes.c_ubyte)
    adv.GetSidSubAuthority.argtypes = [wintypes.LPVOID, wintypes.DWORD]
    adv.GetSidSubAuthority.restype = ctypes.POINTER(wintypes.DWORD)
    process = kernel.OpenProcess(0x1000, False, pid)
    assert process
    token = wintypes.HANDLE()
    try:
        assert adv.OpenProcessToken(process, 8, ctypes.byref(token))
        size = wintypes.DWORD()
        adv.GetTokenInformation(token, 25, None, 0, ctypes.byref(size))
        data = ctypes.create_string_buffer(size.value)
        assert adv.GetTokenInformation(token, 25, data, size, ctypes.byref(size))
        sid = ctypes.cast(data, ctypes.POINTER(ctypes.c_void_p)).contents.value
        count = adv.GetSidSubAuthorityCount(sid).contents.value
        rid = adv.GetSidSubAuthority(sid, count - 1).contents.value
        assert rid == 8192, rid
        return rid
    finally:
        if token:
            kernel.CloseHandle(token)
        kernel.CloseHandle(process)


def identity(pid, executable):
    assert driver.image_path(pid).resolve() == executable.resolve()
    assert driver.process_package_family(pid) == FAMILY
    return {
        "actual_pid": pid,
        "actual_package_family": FAMILY,
        "executable_sha256": digest(executable),
        "integrity_rid": medium_integrity(pid),
    }


def cli_conversion(source, target, output):
    output.mkdir()
    # The registered GUI executable also supports --to. Launching a helper
    # directly from WindowsApps would bypass package activation/identity.
    cli = a.gui
    env = dict(os.environ)
    env["PATH"] = str(Path(os.environ["SystemRoot"]) / "System32")
    for name in (
        "PYTHONPATH",
        "PYTHONHOME",
        "QT_PLUGIN_PATH",
        "QT_QPA_PLATFORM_PLUGIN_PATH",
        "QT_MEDIA_BACKEND",
        "QT_QPA_PLATFORM",
        "QT_SCALE_FACTOR",
        "CITRICA_OFFICE_ENGINE",
        "CITRICA_WINDOWS_ENGINES",
        "CITRICA_TOOL_DIR",
        "CITRICA_OCR_DIR",
    ):
        env.pop(name, None)
    process = subprocess.Popen(
        [str(cli), str(source), "--to", target, "--output", str(output)],
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    details = identity(process.pid, cli)
    stdout, stderr = process.communicate(timeout=180)
    assert process.returncode == 0, (process.returncode, stderr.decode("utf-8", errors="replace")[-1500:])
    files = list(output.glob("*." + target))
    assert len(files) == 1, files
    details.update(exit_code=process.returncode, stderr=stderr.decode("utf-8", errors="replace")[-1500:])
    return files[0], details


try:
    video = a.state / "Movimento.mp4"
    shutil.copy2(Path(r"C:\CitricaVM\v070-store-video-top\Movimento.mp4"), video)
    video_sha = digest(video)
    driver.open("trim", video, "Cortar tempo")
    assert driver.pid == driver.process.pid
    editor_identity = identity(driver.pid, a.gui)
    timeline_name = "Linha do tempo; arraste as alças amarelas para selecionar o trecho"
    driver.wait(lambda rows: any(r["name"] == timeline_name and r["enabled"] for r in rows))
    driver.value("Início", 1)
    driver.value("Duração", 2)
    played = play_and_pause(driver, 1, 3)
    timeline = driver.find(timeline_name)
    dpi = driver.user.GetDpiForWindow(driver.hwnd) / 96.0
    left = timeline["x"] + 14 * dpi
    width = timeline["width"] - 28 * dpi
    y = timeline["y"] + 45 * dpi
    driver.click((left + width * 2.5 / 6, y))
    rows = driver.wait(lambda rows: abs(playback_seconds(rows) - 2.5) < 0.16)
    seek = playback_seconds(rows)
    output = driver.save(video)
    observed = duration(a.oracle_ffprobe, output)
    assert abs(observed - 2) < 0.16
    driver.close()
    assert digest(video) == video_sha
    report["cases"]["video_player_seek_export_normal_exit0"] = {
        **editor_identity,
        "playback_position_seconds": played,
        "actual_seek_position_seconds": seek,
        "output_duration_seconds": observed,
        "normal_editor_exit_code": 0,
        "original_hash_unchanged": True,
    }
    source = a.state / "Editable.docx"
    doc = Document()
    marker = "CITRICA QA 070 MSIX editable text"
    doc.add_paragraph(marker)
    doc.save(source)
    source_sha = digest(source)
    pdf, details = cli_conversion(source, "pdf", a.state / "PDF")
    with pymupdf.open(pdf) as document:
        text = " ".join(page.get_text() for page in document)
    assert marker in " ".join(text.split()), text
    report["cases"]["docx_to_pdf_packaged_registered_executable"] = {
        **details,
        "independent_pdf_text_verified": True,
    }
    pdf_sha = digest(pdf)
    docx, details = cli_conversion(pdf, "docx", a.state / "DOCX")
    converted = Document(docx)
    text = " ".join(p.text for p in converted.paragraphs)
    for table in converted.tables:
        text += " " + " ".join(cell.text for row in table.rows for cell in row.cells)
    assert marker in " ".join(text.split()), text
    assert digest(source) == source_sha and digest(pdf) == pdf_sha
    report["cases"]["pdf_to_editable_docx_packaged_registered_executable"] = {
        **details,
        "editable_document_parsed_and_text_verified": True,
        "original_docx_and_pdf_hash_unchanged": True,
    }
    report["result"] = "passed"
except Exception as error:
    report.update(result="failed", error=str(error))
    raise
finally:
    if driver.process is not None:
        if driver.hwnd:
            driver.user.PostMessageW(driver.hwnd, 0x10, 0, 0)
        try:
            report["failure_normal_cleanup_exit_code"] = driver.process.wait(timeout=30)
        except subprocess.TimeoutExpired:
            report["failure_normal_cleanup_pending"] = True
    (a.state / "report.json").write_text(json.dumps(report, ensure_ascii=True, indent=2), encoding="utf-8")
print(json.dumps({"result": report["result"], "report": str(a.state / "report.json")}))
