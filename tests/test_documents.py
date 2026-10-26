import hashlib
import io
import threading

import pymupdf
import pytest
from docx import Document
from PIL import Image

from citrica.core import Cancelled, ConversionError, convert, formats
from citrica.desktop import request


def test_pdf_docx_text_table_image_and_original(tmp_path):
    source = tmp_path / 'documento com espaços.pdf'
    original = pymupdf.open()
    page = original.new_page()
    page.insert_text((50, 50), 'Citrica editable document', fontsize=18)
    page.insert_text((50, 90), 'First page content')
    image = io.BytesIO()
    Image.new('RGB', (100, 40), 'orange').save(image, format='PNG')
    page.insert_image(pymupdf.Rect(50, 120, 150, 160), stream=image.getvalue())
    for y in [200, 230, 260]:
        page.draw_line((50, y), (350, y))
    for x in [50, 200, 350]:
        page.draw_line((x, 200), (x, 260))
    for x, y, text in [(60, 220, 'Product'), (210, 220, 'Price'), (60, 250, 'Orange'), (210, 250, '12')]:
        page.insert_text((x, y), text)
    page = original.new_page()
    page.insert_text((50, 60), 'Second page retained')
    original.save(source)
    original.close()
    digest = hashlib.sha256(source.read_bytes()).digest()
    assert 'docx' in formats(source)
    output = convert(source, 'docx')[0]
    doc = Document(output)
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Citrica editable document' in text and 'Second page retained' in text
    assert len(doc.inline_shapes) >= 1
    assert doc.tables and doc.tables[0].cell(1, 0).text.strip() == 'Orange'
    assert hashlib.sha256(source.read_bytes()).digest() == digest
    assert convert(source, 'docx')[0] != output
    palette = request({'op': 'palette', 'paths': [str(source)]})
    assert any(i['id'] == 'docx' for i in palette['modes']['convert'])


def test_docx_without_office_keeps_content_in_pdf_and_txt(tmp_path, monkeypatch):
    monkeypatch.setenv('CITRICA_OFFICE_ENGINE', 'builtin')
    source = tmp_path / 'carta.docx'
    doc = Document()
    doc.add_heading('Título da Cítrica', level=1)
    doc.add_paragraph('Conteúdo com acentuação, ç e ã.').runs[0].bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Produto'
    table.cell(0, 1).text = 'Valor'
    table.cell(1, 0).text = 'Laranja'
    table.cell(1, 1).text = '12'
    image = io.BytesIO()
    Image.new('RGB', (100, 40), 'orange').save(image, format='PNG')
    image.seek(0)
    doc.add_picture(image)
    doc.save(source)
    before = source.read_bytes()
    assert formats(source) == ['pdf', 'txt']
    output = convert(source, 'pdf')[0]
    with pymupdf.open(output) as pdf:
        text = ''.join(p.get_text() for p in pdf)
        assert all(s in text for s in ['Título da Cítrica', 'Conteúdo', 'Produto', 'Laranja', '12'])
        assert any(page.get_images() for page in pdf)
    text = convert(source, 'txt')[0].read_text(encoding='utf-8')
    assert 'Conteúdo com acentuação, ç e ã.' in text and 'Laranja\t12' in text
    assert source.read_bytes() == before


def test_docx_txt_builtin_and_cancellation(tmp_path, monkeypatch):
    monkeypatch.setenv('CITRICA_OFFICE_ENGINE', 'builtin')
    text = tmp_path / 'entrada.txt'
    text.write_text('Olá Cítrica\nSegunda linha', encoding='utf-8')
    docx = convert(text, 'docx')[0]
    assert [p.text for p in Document(docx).paragraphs] == ['Olá Cítrica', 'Segunda linha']
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(Cancelled):
        convert(docx, 'pdf', cancel=cancel)
    assert not list(tmp_path.glob('*.pdf'))


def test_pdf_docx_rejects_encrypted_and_corrupt_without_output(tmp_path):
    for name, protected in [('protected.pdf', True), ('corrupt.pdf', False)]:
        source = tmp_path / name
        if protected:
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(source, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='secret')
        else:
            source.write_bytes(b'bad pdf')
        with pytest.raises(ConversionError):
            convert(source, 'docx')
        assert not source.with_suffix('.docx').exists()
    assert not list(tmp_path.glob('.citrica-*'))


def test_pdf_docx_worker_cancel_does_not_publish_partial_file(tmp_path):
    source = tmp_path/'long.pdf'
    with pymupdf.open() as pdf:
        for _ in range(30):
            pdf.new_page().insert_text((50, 60), 'Citrica cancellable conversion')
        pdf.save(source)
    original = source.read_bytes()
    cancel = threading.Event()
    timer = threading.Timer(.1, cancel.set)
    timer.start()
    try:
        with pytest.raises(Cancelled):
            convert(source, 'docx', cancel=cancel)
    finally:
        timer.cancel()
    assert source.read_bytes() == original
    assert not list(tmp_path.glob('*.docx')) and not list(tmp_path.glob('.citrica-*'))
